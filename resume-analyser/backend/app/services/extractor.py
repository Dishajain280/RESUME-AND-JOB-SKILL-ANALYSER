"""
Resume text extraction utilities.
Supports PDF (PyMuPDF), DOCX (python-docx), plain text / Markdown / CSV,
HTML, RTF, ODT, and images (JPG/PNG/… via Tesseract OCR).

Legacy binary .doc (OLE2) is deliberately unsupported — python-docx cannot
parse it, so it is rejected with a clear message instead of a 500.
"""
from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Optional

import structlog

log = structlog.get_logger()

# Plain-text-like formats decoded directly (UTF-8 with replacement).
_PLAIN_TEXT_EXTS = frozenset({".txt", ".md", ".markdown", ".csv"})
# Raster formats OCR'd with Tesseract.
_IMAGE_EXTS = frozenset({
    ".jpg", ".jpeg", ".png", ".gif", ".bmp",
    ".tiff", ".tif", ".webp",
})

# Load .env so TESSERACT_CMD / PATH env vars from the backend .env file
# are visible to _detect_ocr() at import time (pydantic-settings reads .env
# into Settings, but os.environ is what pytesseract / shutil.which consult).
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass


# ── OCR support for scanned (image-based) PDFs ────────────────────────────────
# Tesseract must be installed on the system. On Windows it's typically in
# C:\Program Files\Tesseract-OCR\tesseract.exe; set TESSERACT_CMD in the env
# or add it to PATH. If OCR isn't available, scanned PDFs will still upload
# successfully — they just won't have extractable text.
_OCR_AVAILABLE: bool = False
_TESSERACT_CMD: Optional[str] = None


def _detect_ocr() -> None:
    """Probe for a usable OCR backend (Tesseract via pytesseract)."""
    global _OCR_AVAILABLE, _TESSERACT_CMD
    try:
        import pytesseract  # noqa: F401
        from PIL import Image  # noqa: F401
    except ImportError:
        log.warning("pytesseract or Pillow not installed; OCR disabled")
        return
    binary = os.environ.get("TESSERACT_CMD", "")
    if not binary:
        import shutil
        binary = shutil.which("tesseract")
    if binary:
        _TESSERACT_CMD = binary
        _OCR_AVAILABLE = True
    else:
        log.warning("tesseract binary not found; OCR disabled")


_detect_ocr()


def _ocr_page(page, matrix_dpi: float = 2.0) -> str:
    """OCR a single PDF page using Tesseract (for scanned/image-only PDFs)."""
    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = _TESSERACT_CMD
    pix = page.get_pixmap(matrix=fitz.Matrix(matrix_dpi, matrix_dpi))
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    return pytesseract.image_to_string(img, config="--oem 3 --psm 6")


def validate_magic_bytes(file_bytes: bytes, ext: str) -> Optional[str]:
    """
    Content-sniff the file signature so a renamed executable/photo can't pose
    as a resume. Returns an error message when the content doesn't match the
    declared extension, or None when it's valid.
    """
    ext = ext.lower()
    if ext == ".pdf":
        # PDFs may carry a UTF-8 BOM or tool-specific junk before the header
        # (some scanners/preprocessors prepend bytes); scan the first 4 KB
        # rather than requiring an exact offset-0 match.
        head = file_bytes[:4096]
        if b"%PDF-" not in head:
            return "File content does not look like a valid PDF. Re-export it as a PDF and try again."
    elif ext == ".docx":
        # .docx is a ZIP container (OOXML) — every valid one starts with PK\x03\x04
        if not file_bytes.startswith(b"PK\x03\x04"):
            return "File content is not a valid .docx document (missing ZIP signature)."
        # A real OOXML Word document contains this content-type marker in its
        # package manifest. Scan a generous window (up to 256 KB) to allow for
        # tools that reorder ZIP entries or embed extra metadata at the start.
        if b"[Content_Types].xml" not in file_bytes[:262144]:
            return "File content does not match an OOXML Word document."
    elif ext in _PLAIN_TEXT_EXTS or ext in (".html", ".htm"):
        # Heuristic binary sniff: NUL bytes or a high share of non-printable
        # characters means this isn't human-readable text.
        sample = file_bytes[:8192]
        if b"\x00" in sample:
            return "Text file contains binary data. Upload a plain-text resume."
        if sample:
            non_text = sum(1 for b in sample if b < 0x09 or (0x0E <= b <= 0x1F))
            if non_text / len(sample) > 0.10:
                return "Text file contains binary data. Upload a plain-text resume."
    elif ext == ".rtf":
        if b"{\\rtf" not in file_bytes[:64]:
            return "File content is not a valid RTF document."
    elif ext == ".odt":
        if not file_bytes.startswith(b"PK\x03\x04"):
            return "File content is not a valid .odt document (missing ZIP signature)."
        # The ODF mimetype entry may be stored or deflated, so read it
        # through zipfile instead of scanning raw bytes.
        try:
            import zipfile

            with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
                names = set(z.namelist())
                if "content.xml" not in names:
                    return "File content does not match an OpenDocument Text (.odt) file."
                if "mimetype" in names:
                    mt = z.read("mimetype")[:128]
                    if b"application/vnd.oasis.opendocument.text" not in mt:
                        return "File content does not match an OpenDocument Text (.odt) file."
        except Exception:
            return "File content does not match an OpenDocument Text (.odt) file."
    elif ext in _IMAGE_EXTS:
        sig = file_bytes[:16]
        is_image = (
            sig.startswith(b"\xff\xd8\xff")            # JPEG
            or sig.startswith(b"\x89PNG\r\n\x1a\n")   # PNG
            or sig.startswith(b"GIF8")                   # GIF
            or sig.startswith(b"BM")                     # BMP
            or sig.startswith(b"II*\x00")                # TIFF (little-endian)
            or sig.startswith(b"MM\x00*")                # TIFF (big-endian)
            or (sig.startswith(b"RIFF") and file_bytes[8:12] == b"WEBP")  # WebP
        )
        if not is_image:
            return f"File content does not look like a valid {ext.upper()} image."
    return None


def extract_text(file_bytes: bytes, filename: str) -> tuple[str, int]:
    """
    Return (raw_text, page_count) from file bytes.
    page_count is 1 for DOCX / TXT.
    """
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_pdf(file_bytes)
    elif ext == ".docx":
        return _extract_docx(file_bytes), 1
    elif ext in _PLAIN_TEXT_EXTS:
        return file_bytes.decode("utf-8", errors="replace"), 1
    elif ext in (".html", ".htm"):
        return _extract_html(file_bytes), 1
    elif ext == ".rtf":
        return _extract_rtf(file_bytes), 1
    elif ext == ".odt":
        return _extract_odt(file_bytes), 1
    elif ext in _IMAGE_EXTS:
        return _extract_image(file_bytes), 1
    elif ext == ".doc":
        raise ValueError(
            "Legacy .doc files are not supported. Please re-save your resume "
            "as .docx or PDF and upload again."
        )
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(file_bytes: bytes) -> tuple[str, int]:
    try:
        global fitz  # noqa: PLW0603 — module-level cache for _ocr_page
        try:
            import fitz  # PyMuPDF
        except ImportError:
            global fitz
            log.warning("PyMuPDF not installed, falling back to raw bytes decode")
            return file_bytes.decode("utf-8", errors="replace"), 1
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        pages: list[str] = []
        for page in doc:
            text = page.get_text("text")
            # If the page has no extractable text it's likely a scanned
            # image — fall back to OCR when Tesseract is available.
            if not text.strip() and _OCR_AVAILABLE:
                log.info("pdf_scanned_page_ocr", page=page.number)
                try:
                    text = _ocr_page(page)
                except Exception as ocr_exc:
                    log.warning("ocr_failed", error=str(ocr_exc))
            pages.append(text)
        return "\n".join(pages), len(pages)
    except Exception as exc:
        log.error("pdf_extraction_error", error=str(exc))
        raise


def _extract_html(file_bytes: bytes) -> str:
    """Strip HTML tags, keeping visible text (scripts/styles dropped)."""
    from html.parser import HTMLParser

    class _TextOnly(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self._parts: list[str] = []
            self._skip = 0

        def handle_starttag(self, tag, attrs):
            if tag in ("script", "style", "head"):
                self._skip += 1
            elif tag in ("p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
                self._parts.append("\n")

        def handle_endtag(self, tag):
            if tag in ("script", "style", "head") and self._skip:
                self._skip -= 1
            elif tag in ("p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
                self._parts.append("\n")

        def handle_data(self, data):
            if not self._skip:
                self._parts.append(data)

    parser = _TextOnly()
    parser.feed(file_bytes.decode("utf-8", errors="replace"))
    text = "".join(parser._parts)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _extract_rtf(file_bytes: bytes) -> str:
    """Best-effort RTF → plain text: drop control words, keep text runs."""
    text = file_bytes.decode("utf-8", errors="replace")
    # Escaped characters like \'e9 (Windows-1252 accents)
    text = re.sub(r"\\'[0-9a-fA-F]{2}", "", text)
    # Paragraph / line breaks become real newlines
    text = re.sub(r"\\(par|line|page)\s?", "\n", text)
    # Control words: \word, \wordN, \wordN<space>
    text = re.sub(r"\\[a-z]+-?\d*\s?", " ", text)
    # Escaped braces / backslash
    text = re.sub(r"\\([{}\\])", r"\1", text)
    text = text.replace("{", " ").replace("}", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _extract_odt(file_bytes: bytes) -> str:
    """OpenDocument Text is a ZIP; pull text out of content.xml."""
    import zipfile

    with zipfile.ZipFile(io.BytesIO(file_bytes)) as z:
        names = z.namelist()
        content_name = (
            "content.xml"
            if "content.xml" in names
            else next((n for n in names if n.endswith("content.xml")), None)
        )
        if content_name is None:
            raise ValueError("ODT archive is missing content.xml")
        xml = z.read(content_name).decode("utf-8", errors="replace")
    xml = re.sub(r"<text:line-break[^>]*/>", "\n", xml)
    xml = re.sub(r"</text:p>", "\n", xml)
    xml = re.sub(r"</text:span>", " ", xml)
    return re.sub(r"<[^>]+>", "", xml).strip()


def _extract_image(file_bytes: bytes) -> str:
    """OCR a raster image (screenshot / photo of a resume)."""
    if not _OCR_AVAILABLE:
        raise ValueError(
            "Image files require Tesseract OCR. Install Tesseract and set "
            "TESSERACT_CMD in the backend .env file."
        )
    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = _TESSERACT_CMD
    img = Image.open(io.BytesIO(file_bytes))
    return pytesseract.image_to_string(img, config="--oem 3 --psm 6")


def _extract_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Also grab tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
        return "\n".join(paragraphs)
    except ImportError:
        log.warning("python-docx not installed")
        return file_bytes.decode("utf-8", errors="replace")
    except Exception as exc:
        log.error("docx_extraction_error", error=str(exc))
        raise


# ── Section detection helpers ─────────────────────────────────────────────────

_SECTION_HEADERS = {
    "experience": re.compile(
        r"(work\s+experience|professional\s+experience|employment|career\s+history|experience)",
        re.I,
    ),
    "education": re.compile(
        r"(education|academic|qualifications?|degrees?|university|college)",
        re.I,
    ),
    "skills": re.compile(
        r"(skills?|technical\s+skills?|core\s+competencies|proficiencies|technologies)",
        re.I,
    ),
    "summary": re.compile(
        r"(summary|profile|objective|about\s+me|overview)",
        re.I,
    ),
    "certifications": re.compile(
        r"(certifications?|certificates?|licenses?|credentials?|courses?)",
        re.I,
    ),
}


def detect_sections(text: str) -> dict[str, str]:
    """
    Very lightweight section splitter — returns a dict mapping section names
    to their raw text content.
    """
    lines = text.split("\n")
    sections: dict[str, list[str]] = {"header": []}
    current = "header"

    for line in lines:
        stripped = line.strip()
        matched = False
        for name, pattern in _SECTION_HEADERS.items():
            if pattern.fullmatch(stripped) or (
                len(stripped) < 40 and pattern.search(stripped)
            ):
                current = name
                sections.setdefault(current, [])
                matched = True
                break
        if not matched:
            sections.setdefault(current, []).append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if v}
